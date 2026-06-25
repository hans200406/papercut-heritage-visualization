from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "data_cleaning_process.docx"


def set_font(run, size=None, color=None, bold=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_para(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_font(r, 11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r, 10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_font(r, 10.5)
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.85)
sec.bottom_margin = Inches(0.85)
sec.left_margin = Inches(0.9)
sec.right_margin = Inches(0.9)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.25
normal.paragraph_format.space_after = Pt(7)

for style_name, size, color in [
    ("Heading 1", 15, "B21F2D"),
    ("Heading 2", 12.5, "7D1420"),
]:
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = True

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("中国剪纸非遗数据来源与清洗过程说明")
set_font(r, 20, "B21F2D", True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("纸上生花：中国剪纸非遗多维信息可视化设计")
set_font(r, 12, "755F55", True)

doc.add_heading("一、文件说明", level=1)
add_para(
    doc,
    "本说明用于配合课程设计提交中的“原始数据集及清洗过程说明”。项目数据围绕中国剪纸类国家级非物质文化遗产代表性项目整理，共形成54条记录，并分别保留原始整理表、清洗后数据表和前端使用的数据文件。",
)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["文件名", "类型", "用途"]
for i, text in enumerate(headers):
    table.rows[0].cells[i].text = text
rows = [
    ("raw_heritage_papercut_projects.csv", "原始整理表/来源留痕表", "保留公开名录核心字段，并记录来源说明、清洗动作、坐标处理和关键词处理说明。"),
    ("heritage_papercut_projects.csv", "清洗后数据集", "用于模块A图表分析和模块B生成艺术的数据基础。"),
    ("papercutData.js", "前端数据文件", "由清洗后CSV转换而来，供网页直接读取。"),
]
for row in rows:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        cells[i].text = text

for row in table.rows:
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(2)
            for run in para.runs:
                set_font(run, 9.5)

doc.add_heading("二、数据来源", level=1)
add_bullet(doc, "中国非物质文化遗产网国家级非物质文化遗产代表性项目名录公开资料：用于整理剪纸类项目名称、申报地区、批次和类别。")
add_bullet(doc, "UNESCO Intangible Cultural Heritage 中关于 Chinese paper-cut 的项目介绍：用于确认中国剪纸的国际代表性背景。")
add_bullet(doc, "公开行政区与地理信息资料：用于按申报地或代表性传承地补充近似经纬度，服务网页空间可视化表达。")

doc.add_heading("三、字段说明", level=1)
field_table = doc.add_table(rows=1, cols=3)
field_table.alignment = WD_TABLE_ALIGNMENT.CENTER
field_table.style = "Table Grid"
for i, text in enumerate(["字段", "含义", "使用位置"]):
    field_table.rows[0].cells[i].text = text
field_rows = [
    ("项目名称", "剪纸类非遗项目名称", "项目识别与标题展示"),
    ("省级地区", "项目所在省、自治区或直辖市", "省份分布统计"),
    ("地级/县级地区", "项目申报或代表性传承地区", "原始来源留痕与地区说明"),
    ("批次", "进入国家级非遗名录或扩展名录的批次", "批次结构环形图"),
    ("类别", "统一为传统美术", "数据筛选依据"),
    ("经度/纬度", "申报地或传承地近似坐标", "地图或空间可视化备用字段"),
    ("视觉关键词", "从地域、民族、纹样和工艺特征提炼的图形线索", "关键词频次统计与p5.js生成艺术"),
]
for row in field_rows:
    cells = field_table.add_row().cells
    for i, text in enumerate(row):
        cells[i].text = text
for row in field_table.rows:
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(2)
            for run in para.runs:
                set_font(run, 9.5)

doc.add_heading("四、清洗过程", level=1)
steps = [
    "从公开名录中筛选名称或项目说明中含“剪纸、刻纸、凿花、过门笺”等关键词的传统美术类项目。",
    "将申报地区拆分为省级地区与地级/县级地区，便于后续按省份统计和展示。",
    "将跨地区或综合类项目进行合并处理，例如保留其主要代表性地区和综合项目名称，避免重复计数。",
    "统一批次命名为“第一批、第二批扩展、第三批扩展、第四批扩展、第五批扩展”。",
    "将项目类别统一标注为“传统美术”，保证数据口径一致。",
    "按申报地或代表性传承地补充近似经纬度，用于信息可视化表达，不作为精确测绘坐标。",
    "从项目名称、民族属性、地域民俗和纹样特征中提炼视觉关键词，用于连接模块A统计分析和模块B剪纸生成艺术。",
    "将清洗后的CSV转换为papercutData.js，使网页可以在无后端环境下直接读取数据。",
]
for step in steps:
    add_number(doc, step)

doc.add_heading("五、数据限制与使用说明", level=1)
add_bullet(doc, "项目名称、申报地区、批次和类别来自公开名录整理，适合作为课程设计的数据基础。")
add_bullet(doc, "经纬度为近似定位，仅用于视觉表达和空间感知，不用于精确地图测绘或行政边界判断。")
add_bullet(doc, "视觉关键词属于设计性编码字段，目的是把数据分析结果转译为生成艺术元素，因此在报告中作为“视觉提取依据”使用。")
add_bullet(doc, "清洗后数据主要服务本课程设计的模块A图表、模块B生成艺术和模块C叙事整合。")

doc.add_heading("六、初步数据洞察", level=1)
insights = [
    "剪纸类非遗项目具有明显的地域集中特征，辽宁、山西、江苏、福建、山东等省份数量较多。",
    "第二批扩展名录中的剪纸类项目数量最多，说明剪纸类项目在国家级非遗保护体系中经历了集中扩充和细分认定。",
    "剪纸项目不仅包含汉族民间窗花，也包含满族、苗族、水族、回族、傣族等民族文化样式。",
    "多数项目的视觉关键词集中在“窗花、花鸟、吉祥、民俗、几何纹、民族纹样”等方向，为生成艺术提供图形元素依据。",
]
for insight in insights:
    add_bullet(doc, insight)

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = ""
    r = footer.add_run("纸上生花：中国剪纸非遗多维信息可视化设计")
    set_font(r, 9, "755F55")

doc.save(OUT)
print(OUT)
