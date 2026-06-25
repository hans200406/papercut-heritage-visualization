# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from pathlib import Path

root = Path.cwd()
out = root / 'design_report_papercut.docx'
img_a = root / 'screenshots' / 'module_a_data_analysis.png'
img_b = root / 'screenshots' / 'module_b_papercut_art_focus.png'
img_c = root / 'screenshots' / 'module_c_story_focus.png'
img_process = root / 'screenshots' / 'process_sketch.png'

sections = [
('一、项目背景与选题意义', [
'本课程设计以“中国剪纸非遗”为研究与创作对象，作品题名为“纸上生花”。剪纸是中国民间美术中传播范围广、参与群体多、节庆属性强的非物质文化遗产形式。它既是一种手工技艺，也是一套关于地域、民俗、信仰、审美和生活愿望的视觉语言。窗花、花鸟、生命树、吉祥纹样、民族图腾等元素，在不同地区形成了差异明显的风格谱系。',
'课程设计任务要求同时完成“理性数据分析”和“感性图形艺术”两个核心模块，因此本项目没有把剪纸仅作为装饰素材处理，而是尝试将其转化为一个可阅读、可交互、可传播的数字媒体作品。模块A通过真实数据呈现剪纸类非遗项目的地区分布、名录批次和视觉关键词；模块B通过 p5.js 生成艺术重构剪纸的对称、镂空和红色民俗意象；模块C再通过滚动叙事将两者组织为一个完整网页体验。',
'选题的意义主要体现在三个方面：第一，剪纸具有鲜明的视觉识别度，适合转化为数字生成图形；第二，剪纸类非遗项目在国家级名录中数量较多，便于形成结构化数据分析；第三，剪纸与节庆、婚俗、民族文化和地域民俗高度相关，能够支持从宏观分布到微观纹样的多层次叙事。'
]),
('二、数据来源与处理方法', [
'本项目整理了国家级非物质文化遗产代表性项目名录中与“剪纸、刻纸、凿花、过门笺”等相关的传统美术项目，共形成54条记录。数据字段包括项目名称、省级地区、地级或县级地区、批次、类别、经度、纬度和视觉关键词。数据来源主要包括国家级非遗名录公开资料、UNESCO对Chinese paper-cut的项目介绍，以及各项目申报地的公开地理信息。',
'原始资料以项目名称和申报地区为核心，经过人工筛选和结构化整理后，先形成raw_heritage_papercut_projects.csv原始整理表，再转换为网页使用的heritage_papercut_projects.csv清洗数据表。原始整理表保留来源类型、来源说明、清洗动作、坐标处理说明和关键词处理说明，用于回应课程设计中“原始数据集及清洗过程说明”的提交要求。',
'清洗过程包括：筛选剪纸相关项目，统一“第一批、第二批扩展、第三批扩展、第四批扩展、第五批扩展”等批次口径，将项目类别统一为“传统美术”，并为每条记录补充近似经纬度，方便后续进行地理可视化或空间散点表达。经纬度仅作为信息可视化中的近似定位依据，不作为精确测绘坐标。',
'视觉关键词字段是连接模块A和模块B的重要桥梁。该字段不是简单的统计字段，而是从项目名称、民族属性、地域特征和工艺特点中提炼出的图形线索，例如“窗花、花鸟、吉祥、几何纹、满族纹样、苗族纹样、孔雀、生命树”等。这些关键词既用于模块A的关键词频次分析，也用于模块B的生成艺术设计。'
]),
('三、模块A：数据分析可视化设计', [
'模块A的目标是建立剪纸非遗的“宏观画像”。本项目没有选择单一图表堆砌，而是围绕“地域分布—批次结构—视觉谱系”建立三层分析路径。第一张图表为省级地区项目数量柱状图，用于比较不同省份的项目集中程度；第二张图表为名录批次结构环形图，用于呈现剪纸项目被纳入保护体系的时间结构；第三张图表为视觉关键词频次图，用于分析剪纸项目背后反复出现的图形与文化元素。',
'省级分布图采用横向柱状图，是因为省份名称较多，横向排布可以减少标签拥挤，同时便于按数量排序。批次结构采用环形图，是因为各批次之间构成整体的一部分，适合展示比例关系。视觉关键词采用条形图，是因为关键词之间没有连续时间关系，更适合用频次比较展示其重要程度。',
'从数据中可以得到三条主要洞察：第一，剪纸类非遗项目具有明显的地域集中特征，辽宁、山西、江苏、福建、山东等地数量较多，说明剪纸传承与地方民俗、节庆活动和区域手工艺生态密切相关。第二，第二批扩展名录中的剪纸类项目数量最多，体现出剪纸项目在国家级非遗保护体系中经历了集中补充和细分认定。第三，剪纸并不只是汉族窗花样式，还包括满族、苗族、水族、回族、傣族等多民族剪纸形态，表现出中国剪纸的多元文化结构。'
]),
('四、模块B：图形艺术信息可视化设计', [
'模块B围绕剪纸的“微观美学”展开。课程任务要求不能直接堆砌照片，因此本项目采用代码生成艺术方式，将剪纸的核心视觉规律转化为p5.js图形系统。生成器主要提取了三类视觉特征：对称结构、镂空负形和红色民俗色彩。对称结构来自团花、窗花等传统剪纸构图；镂空负形来自剪纸的“剪去部分形成图案”的工艺逻辑；红色主色来自中国民俗中喜庆、祈福和节庆的色彩传统。',
'生成艺术模块提供三种纹样模式：团花、花鸟和吉祥纹。团花模式强调中心放射和重复花瓣结构，表现窗花式的秩序感；花鸟模式抽象出鸟形、翅膀和花瓣，呼应剪纸中常见的花鸟题材；吉祥纹模式使用方胜、圆形和几何环形结构，表现传统吉祥图案的符号化特征。用户可以通过按钮切换纹样类型，也可以通过滑杆调整复杂度，使图形在简洁和繁复之间变化。',
'该模块并不是孤立的装饰部分，而是从模块A的数据关键词中延展出来的视觉转译。例如，数据中高频出现的“窗花、花鸟、吉祥、几何纹、民族纹样”等词汇，分别对应生成器中的对称花形、花鸟图形和几何吉祥纹。这样，数据分析结果与艺术生成逻辑之间形成了清晰的因果关系。'
]),
('五、模块C：文化叙事与交互设计', [
'模块C负责将理性数据和感性图形整合为一个完整的数字媒体作品。页面结构遵循“吸引—解释—分析—转译—总结”的叙事顺序：开场以红色剪纸符号和“纸上生花”标题建立文化氛围；随后进入模块A，用真实数据解释剪纸非遗的分布与保护现状；中段设置“叙事逻辑”过渡区，说明如何从数据中的地域和关键词进入纹样生成；后段进入模块B，让观众通过交互操作感受剪纸图形的生成过程；最后以“数字技术让剪纸被看见”总结作品意义。',
'交互设计包含三个层面。第一，ECharts图表支持悬停提示，用户可以查看各省份、批次和关键词的具体数量。第二，页面具有固定导航，用户可以快速跳转到开场、数据分析、叙事逻辑、生成艺术和总结区。第三，p5.js生成艺术模块提供按钮和滑杆交互，用户可以切换团花、花鸟、吉祥纹，并调整纹样复杂度。',
'视觉风格采用“新中式数字展览”方向。整体配色以剪纸红、深红、宣纸米白、墨色和克制金色为主；布局采用大留白、纸质面板和简洁卡片，使页面既有传统文化气息，又符合现代网页的阅读习惯。图表区域保持清晰理性，生成艺术区域则强化审美表达，从而形成理性和感性的互补。'
]),
('六、作品实现与文件结构', [
'作品采用纯前端方式实现，便于本地直接运行，也便于后续部署到GitHub Pages或Vercel。页面主体文件为index.html，样式文件为styles/main.css，模块A图表逻辑为scripts/app.js，模块B生成艺术逻辑为scripts/papercut-art.js，数据文件包括data/heritage_papercut_projects.csv和data/papercutData.js。ECharts用于统计图表绘制，p5.js用于生成艺术绘制。'
]),
('七、总结', [
'“纸上生花”尝试将非遗剪纸从静态展示转化为可分析、可交互、可再生成的数字文化体验。模块A通过真实数据呈现剪纸类非遗的保护格局，帮助观众理解其地域分布、批次变化和文化关键词；模块B通过生成艺术将剪纸的图形规律转化为动态视觉语言，使观众直观感受剪纸的对称、镂空和红色民俗审美；模块C则通过滚动叙事把数据与艺术连接起来，形成完整的观看路径。',
'本项目的核心价值在于把“看见传统”与“理解传统”结合起来。数据图表让非遗保护不再停留在抽象口号中，生成艺术让非遗美学不再只是图片展示。通过信息可视化的方法，剪纸被重新组织为一种兼具知识性、审美性和传播性的数字媒介。'
])]

def set_font(run, size=None, color=None, bold=None):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold

def add_para(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    set_font(r)
    return p

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 9, '755F55')

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.333
normal.paragraph_format.space_after = Pt(8)
for style_name, size, color, before, after in [('Heading 1',16,'2E74B5',18,10),('Heading 2',13,'2E74B5',12,6),('Heading 3',12,'1F4D78',8,4)]:
    s = styles[style_name]
    s.font.name = 'Calibri'
    s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('《信息可视化》课程设计说明书')
set_font(r, 22, 'B21F2D', True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('纸上生花：中国剪纸非遗多维信息可视化设计')
set_font(r, 16, '251714', True)

meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = 'Table Grid'
for i, (k, v) in enumerate([('课程名称','信息可视化'),('课设主题','数字遗珍——非遗文化多维信息可视化设计'),('作品形式','交互网页 / H5 数字媒体项目'),('选题对象','中国剪纸非遗')]):
    meta.cell(i,0).text = k
    meta.cell(i,1).text = v
    for cell in meta.rows[i].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(0)

for title, paras in sections:
    doc.add_heading(title, level=1)
    for text in paras:
        add_para(doc, text)
    if title.startswith('二、'):
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'
        for idx, h in enumerate(['字段','含义','用途']): table.rows[0].cells[idx].text = h
        for row in [('项目名称','剪纸类非遗项目名称','识别具体项目'),('省级地区','项目所在省、自治区或直辖市','省份分布统计'),('批次','进入名录或扩展名录的批次','观察保护进程'),('经度/纬度','项目所在地近似坐标','后续地图或空间展示'),('视觉关键词','纹样、民族、地域和工艺特征','连接数据分析与生成艺术')]:
            cells = table.add_row().cells
            for i, val in enumerate(row): cells[i].text = val
        if img_process.exists():
            doc.add_picture(str(img_process), width=Inches(6.2)); add_caption(doc, '图1  课程设计过程草图：原始资料、数据清洗、模块A、模块B与模块C的关系')
    if title.startswith('三、') and img_a.exists():
        doc.add_picture(str(img_a), width=Inches(6.2)); add_caption(doc, '图2  模块A数据分析页面截图')
    if title.startswith('四、') and img_b.exists():
        doc.add_picture(str(img_b), width=Inches(6.2)); add_caption(doc, '图3  模块B剪纸生成艺术页面截图')
    if title.startswith('五、') and img_c.exists():
        doc.add_picture(str(img_c), width=Inches(6.2)); add_caption(doc, '图4  模块C滚动叙事整合页面截图')
    if title.startswith('六、'):
        for item in ['index.html：网页主体结构，包含开场、模块A、叙事过渡、模块B和结尾。','styles/main.css：新中式剪纸视觉风格、图表面板、导航和生成艺术区域样式。','scripts/app.js：读取剪纸项目数据，生成省份分布、批次结构和关键词频次图。','scripts/papercut-art.js：p5.js剪纸纹样生成器，支持图案切换和复杂度控制。','data/raw_heritage_papercut_projects.csv：原始整理表和来源留痕表。','data/raw_source_notes.md：原始资料范围、清洗转换关系与数据限制说明。','data/heritage_papercut_projects.csv：清洗后的剪纸非遗数据集。','screenshots/process_sketch.png：课程设计过程草图。']:
            p = doc.add_paragraph(style='List Bullet')
            r = p.add_run(item)
            set_font(r)

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = ''
    r = footer.add_run('纸上生花：中国剪纸非遗多维信息可视化设计')
    set_font(r, 9, '755F55')

doc.save(out)
print(out)
